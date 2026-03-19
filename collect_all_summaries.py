import os
import pandas as pd
from openpyxl import load_workbook

OUT_DIR = os.path.join(os.path.dirname(__file__), 'out')
RESULT_XLSX = os.path.join(OUT_DIR, 'final_report_ALL.xlsx')

summary_dfs = []
summary_names = []

for folder in os.listdir(OUT_DIR):
    folder_path = os.path.join(OUT_DIR, folder)
    if not os.path.isdir(folder_path):
        continue
    # Пропустить служебные папки
    if folder.startswith('_'):
        continue
    # Найти xlsx файл
    for file in os.listdir(folder_path):
        if file.endswith('.xlsx') and file.startswith('final_report'):
            file_path = os.path.join(folder_path, file)
            try:
                wb = load_workbook(file_path, read_only=True, data_only=True)
                if 'Summary' in wb.sheetnames:
                    df = pd.read_excel(file_path, sheet_name='Summary')
                    # Имя листа без префикса out_
                    sheet_name = folder.replace('out_', '')
                    summary_dfs.append((sheet_name, df))
                    summary_names.append(sheet_name)
            except Exception as e:
                print(f"Ошибка чтения {file_path}: {e}")


def aggregate_by_email(df):
    """Группирует по Email, суммирует числа, берёт наиболее частое имя автора."""
    if 'Email' not in df.columns or df['Email'].isna().all():
        return df.groupby('Author', as_index=False).sum(numeric_only=True)
    author_by_email = (
        df.groupby('Email')['Author']
        .agg(lambda x: x.value_counts().index[0])
        .reset_index()
    )
    numeric_cols = df.select_dtypes(include='number').columns.tolist()
    agg = df.groupby('Email', as_index=False)[numeric_cols].sum()
    result = author_by_email.merge(agg, on='Email')
    cols = ['Author', 'Email'] + [c for c in result.columns if c not in ('Author', 'Email')]
    return result[cols]

try:
    with pd.ExcelWriter(RESULT_XLSX, engine='openpyxl') as writer:
        # Сначала общий Summary
        if summary_dfs:
            all_df = pd.concat([df for _, df in summary_dfs], ignore_index=True)
            summary_grouped = aggregate_by_email(all_df)
            summary_grouped.to_excel(writer, sheet_name='Summary', index=False)
        else:
            pd.DataFrame({'Нет данных': []}).to_excel(writer, sheet_name='Summary', index=False)
        # Затем остальные листы
        for name, df in summary_dfs:
            df.to_excel(writer, sheet_name=name, index=False)
    print(f"✅ Итоговый отчёт создан: {RESULT_XLSX}")
except PermissionError:
    print(f"❌ Ошибка: Нет доступа к файлу {RESULT_XLSX}. Возможно, файл открыт в Excel. Закройте его и повторите попытку.")
